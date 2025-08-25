from docx import Document

# Create a new Document
doc = Document()

# Add a title
doc.add_heading('Test Document with Numbered Lists', 0)

# Add some paragraphs
doc.add_paragraph('This document contains various types of lists to test DOCX processing capabilities.')

# Add a heading for numbered lists
doc.add_heading('Sample Numbered List', 1)

# Add a numbered list
doc.add_paragraph('First numbered item', style='List Number')
doc.add_paragraph('Second numbered item', style='List Number')
doc.add_paragraph('Third numbered item', style='List Number')

# Add another heading
doc.add_heading('Sample Bulleted List', 1)

# Add a bulleted list
doc.add_paragraph('First bulleted item', style='List Bullet')
doc.add_paragraph('Second bulleted item', style='List Bullet')
doc.add_paragraph('Third bulleted item', style='List Bullet')

# Add a heading for nested lists
doc.add_heading('Nested List Example', 1)

# Add a numbered list with nested items
doc.add_paragraph('Main item 1', style='List Number')
doc.add_paragraph('Sub-item 1.1', style='List Number 2')
doc.add_paragraph('Sub-item 1.2', style='List Number 2')
doc.add_paragraph('Main item 2', style='List Number')
doc.add_paragraph('Sub-item 2.1', style='List Number 2')
doc.add_paragraph('Main item 3', style='List Number')

# Save the document
doc.save('test-numbered-lists.docx')

print("Test DOCX file with numbered lists created successfully: test-numbered-lists.docx")