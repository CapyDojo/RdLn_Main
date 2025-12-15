from docx import Document
from docx.shared import Inches

# Create a new Document
doc = Document()

# Add a title
doc.add_heading('Test Document with Track Changes', 0)

# Add some paragraphs
doc.add_paragraph('This is a test document created to evaluate DOCX processing capabilities.')
doc.add_paragraph('It contains various types of content to test text extraction and track change detection.')

# Add a paragraph with formatting
para = doc.add_paragraph('This paragraph contains ')
para.add_run('bold text').bold = True
para.add_run(' and ')
para.add_run('italic text').italic = True
para.add_run('.')

# Add a heading
doc.add_heading('Sample Content', 1)

# Add a bulleted list
doc.add_paragraph('First item in list', style='List Bullet')
doc.add_paragraph('Second item in list', style='List Bullet')
doc.add_paragraph('Third item in list', style='List Bullet')

# Add a numbered list
doc.add_paragraph('First numbered item', style='List Number')
doc.add_paragraph('Second numbered item', style='List Number')
doc.add_paragraph('Third numbered item', style='List Number')

# Add a table
table = doc.add_table(rows=3, cols=2)
table.style = 'Table Grid'
hdr_cells = table.rows[0].cells
hdr_cells[0].text = 'Column 1'
hdr_cells[1].text = 'Column 2'
row_cells = table.rows[1].cells
row_cells[0].text = 'Cell 1'
row_cells[1].text = 'Cell 2'
row_cells = table.rows[2].cells
row_cells[0].text = 'Cell 3'
row_cells[1].text = 'Cell 4'

# Save the document
doc.save('test-document-with-changes.docx')

print("Test DOCX file created successfully: test-document-with-changes.docx")
print("Note: This document doesn't have track changes yet. You'll need to manually add them in Word.")