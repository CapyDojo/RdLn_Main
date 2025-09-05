from docx import Document

# Create a new Document
doc = Document()

# Add a title
doc.add_heading('Complex List Formatting Test Document', 0)

# Add introduction
doc.add_paragraph('This document contains various complex list formats to test DOCX processing capabilities.')

# Add a heading for standard numbered lists
doc.add_heading('Standard Numbered Lists', 1)

# Add a standard numbered list
doc.add_paragraph('First numbered item', style='List Number')
doc.add_paragraph('Second numbered item', style='List Number')
doc.add_paragraph('Third numbered item', style='List Number')

# Add a heading for bulleted lists
doc.add_heading('Standard Bulleted Lists', 1)

# Add a standard bulleted list
doc.add_paragraph('First bulleted item', style='List Bullet')
doc.add_paragraph('Second bulleted item', style='List Bullet')
doc.add_paragraph('Third bulleted item', style='List Bullet')

# Add a heading for nested lists
doc.add_heading('Nested List Example', 1)

# Add a nested list structure
doc.add_paragraph('Main item 1', style='List Number')
doc.add_paragraph('Sub-item 1.1', style='List Number 2')
doc.add_paragraph('Sub-item 1.2', style='List Number 2')
doc.add_paragraph('Main item 2', style='List Number')
doc.add_paragraph('Sub-item 2.1', style='List Number 2')
doc.add_paragraph('Sub-sub-item 2.1.1', style='List Number 3')
doc.add_paragraph('Sub-sub-item 2.1.2', style='List Number 3')
doc.add_paragraph('Main item 3', style='List Number')

# Add regular paragraphs to test structure preservation
doc.add_paragraph('This is a regular paragraph to test structure preservation.')
doc.add_paragraph('Another regular paragraph with some text to see how line breaks are handled.')

# Add a heading for mixed content
doc.add_heading('Mixed Content with Lists', 1)

doc.add_paragraph('This paragraph comes before a list.')
doc.add_paragraph('1. Numbered list item', style='List Number')
doc.add_paragraph('2. Another numbered list item', style='List Number')
doc.add_paragraph('This paragraph comes after a list.')
doc.add_paragraph('• Bulleted list item', style='List Bullet')
doc.add_paragraph('• Another bulleted list item', style='List Bullet')
doc.add_paragraph('Final paragraph in this section.')

# Save the document
doc.save('complex-lists-test.docx')

print("Complex test DOCX file with various list formats created successfully: complex-lists-test.docx")