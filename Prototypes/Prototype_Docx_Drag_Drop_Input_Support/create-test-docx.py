# create-test-docx.py
# Script to create a simple DOCX file for testing

from docx import Document

# Create a new Document
doc = Document()

# Add a title
doc.add_heading('Test Document for DOCX Processing', 0)

# Add some paragraphs
doc.add_paragraph('This is a simple test document for testing DOCX processing capabilities.')
doc.add_paragraph('It contains several paragraphs of text.')
doc.add_paragraph('This document can be used to verify that text extraction is working correctly.')

# Add some formatted text
para = doc.add_paragraph('This paragraph contains ')
para.add_run('bold text').bold = True
para.add_run(' and ')
para.add_run('italic text').italic = True
para.add_run('.')

# Add a heading
doc.add_heading('Sample List', 1)

# Add a bulleted list
doc.add_paragraph('First item in list', style='List Bullet')
doc.add_paragraph('Second item in list', style='List Bullet')
doc.add_paragraph('Third item in list', style='List Bullet')

# Add a numbered list
doc.add_paragraph('First numbered item', style='List Number')
doc.add_paragraph('Second numbered item', style='List Number')
doc.add_paragraph('Third numbered item', style='List Number')

# Add a table
table = doc.add_table(rows=3, cols=2)
table.style = 'Table Grid'
hdr_cells = table.rows[0].cells
hdr_cells[0].text = 'Column 1'
hdr_cells[1].text = 'Column 2'
row_cells = table.rows[1].cells
row_cells[0].text = 'Cell 1'
row_cells[1].text = 'Cell 2'
row_cells = table.rows[2].cells
row_cells[0].text = 'Cell 3'
row_cells[1].text = 'Cell 4'

# Save the document
doc.save('Prototypes/test-document.docx')

print("Test DOCX file created successfully: Prototypes/test-document.docx")